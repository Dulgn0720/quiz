# -*- coding: utf-8 -*-
"""Байгууллагын дотоод журмын Word баримт үүсгэгч.

Монгол Улсын байгууллагын дотоод эрх зүйн баримт бичгийн нийтлэг хэлбэрийг
баримтална: батламжийн хаяг, хяналтын хүснэгт, "Нэг. …" хэлбэрийн бүлэг,
"1.1." / "1.1.1." хэлбэрийн заалтын дугаарлалт, төгсгөлийн тэмдэг.
"""

import os
import re
import shutil
import zipfile
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

COMPANY_FULL = '«Агула Даатгал» ХХК'
COMPANY_SHORT = "Компани"
DEPARTMENT = "Инноваци технологийн хэлтэс"
HEAD = "Инноваци технологийн хэлтсийн ахлах"
DIRECTOR = "Инноваци технологийн хэлтсийн захирал"
CEO = "Гүйцэтгэх захирал"

OUT_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
FONT = "Times New Roman"

ORDINALS = [
    "Нэг", "Хоёр", "Гурав", "Дөрөв", "Тав", "Зургаа", "Долоо", "Найм",
    "Ес", "Арав", "Арван нэг", "Арван хоёр", "Арван гурав", "Арван дөрөв",
    "Арван тав", "Арван зургаа", "Арван долоо", "Арван найм", "Арван ес",
    "Хорь", "Хорин нэг", "Хорин хоёр", "Хорин гурав", "Хорин дөрөв",
    "Хорин тав", "Хорин зургаа",
]


def _set_font(run, size=12, bold=False, italic=False):
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), FONT)


def _para(doc, text="", size=12, bold=False, italic=False,
          align=WD_ALIGN_PARAGRAPH.JUSTIFY, after=6, before=0,
          indent=None, first_line=None, line=1.15):
    p = doc.add_paragraph()
    p.alignment = align
    pf = p.paragraph_format
    pf.space_after = Pt(after)
    pf.space_before = Pt(before)
    pf.line_spacing = line
    if indent is not None:
        pf.left_indent = Cm(indent)
    if first_line is not None:
        pf.first_line_indent = Cm(first_line)
    if text:
        _set_font(p.add_run(text), size=size, bold=bold, italic=italic)
    return p


def build(meta, chapters, filename):
    """meta   — dict(no, title, owner, approver, review, version)
    chapters  — [(бүлгийн нэр, [заалт | ("sub", n, текст) | ("note", текст)])]
    """
    doc = Document()

    sec = doc.sections[0]
    sec.top_margin = Cm(2.0)
    sec.bottom_margin = Cm(2.0)
    sec.left_margin = Cm(3.0)
    sec.right_margin = Cm(1.5)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(12)

    # ── Батламжийн хаяг ──
    for line in [
        f"{COMPANY_FULL}-ийн Гүйцэтгэх захирлын",
        "20___ оны ___ дугаар сарын ___-ны өдрийн",
        "А/___ дугаар тушаалын хавсралт",
    ]:
        p = _para(doc, line, size=11, align=WD_ALIGN_PARAGRAPH.RIGHT,
                  after=0, line=1.0)
        p.paragraph_format.left_indent = Cm(8.0)

    _para(doc, "", after=16)

    # ── Гарчиг ──
    _para(doc, meta["title"].upper(), size=14, bold=True,
          align=WD_ALIGN_PARAGRAPH.CENTER, after=4, line=1.0)
    _para(doc, f"Баримтын дугаар: {meta['no']}", size=11,
          align=WD_ALIGN_PARAGRAPH.CENTER, after=14, line=1.0)

    # ── Хяналтын хүснэгт ──
    # Гурван шатлал: боловсруулах — хянан баталгаажуулах — батлах.
    # Баримт боловсруулсан хүн өөрөө батлахгүй.
    rows = [
        ("Баримтын төрөл", meta.get("kind", "Журам")),
        ("Баримтын дугаар", meta["no"]),
        ("Хувилбар", meta.get("version", "1.0")),
        ("Боловсруулсан нэгж", DEPARTMENT),
        ("Боловсруулсан албан тушаалтан", meta.get("author", HEAD)),
        ("Хянан баталгаажуулсан", meta.get("reviewer", DIRECTOR)),
        ("Батлах эрх бүхий албан тушаалтан", meta["approver"]),
        ("Хариуцан хэрэгжүүлэгч", meta["owner"]),
        ("Хүчин төгөлдөр болох огноо", "20___ оны ___ сарын ___-ны өдөр"),
        ("Хянан шинэчлэх давтамж", meta["review"]),
    ]
    t = doc.add_table(rows=len(rows), cols=2)
    t.style = "Table Grid"
    for i, (k, v) in enumerate(rows):
        for j, val in enumerate((k, v)):
            cell = t.cell(i, j)
            cell.width = Cm(6.8 if j == 0 else 8.7)
            cp = cell.paragraphs[0]
            cp.paragraph_format.space_after = Pt(2)
            cp.paragraph_format.space_before = Pt(2)
            cp.paragraph_format.line_spacing = 1.0
            _set_font(cp.add_run(val), size=10, bold=(j == 0))

    _para(doc, "", after=12)

    # ── Бүлгүүд ──
    for ci, (chapter_title, clauses) in enumerate(chapters, start=1):
        # Гарчиг хуудасны төгсгөлд ганцаараа үлдэхээс сэргийлнэ.
        head = _para(doc, f"{ORDINALS[ci - 1]}. {chapter_title}", size=12,
                     bold=True, align=WD_ALIGN_PARAGRAPH.CENTER,
                     before=12, after=8, line=1.0)
        head.paragraph_format.keep_with_next = True

        n = 0
        for clause in clauses:
            if isinstance(clause, tuple) and clause[0] == "sub":
                _para(doc, f"{ci}.{n}.{clause[1]}. {clause[2]}",
                      indent=1.0, after=4)
                continue
            if isinstance(clause, tuple) and clause[0] == "note":
                _para(doc, clause[1], italic=True, size=11,
                      indent=1.0, after=6)
                continue
            n += 1
            par = _para(doc, f"{ci}.{n}. {clause}", after=5, first_line=0.0)
            # Хоёр цэгээр төгссөн заалт нь дэд заалтуудынхаа удиртгал тул
            # тэднээсээ хуудсаар тасрах ёсгүй.
            if clause.rstrip().endswith(":"):
                par.paragraph_format.keep_with_next = True

    _para(doc, "", after=10)
    end = _para(doc, "--- о0о ---", align=WD_ALIGN_PARAGRAPH.CENTER,
                size=11, after=0, line=1.0)
    end.paragraph_format.keep_with_next = True

    # ── Гарын үсгийн хэсэг ──
    # Хяналтын хүснэгт нь боловсруулсан, хянасан хоёр албан тушаалыг
    # нэрлэдэг боловч тэдгээрийн гарын үсэг зурах газар байсангүй.
    # Гарын үсэггүй бол «хянан баталгаажуулсан» гэдэг нь батлагдахгүй
    # мэдэгдэл хэвээр үлдэнэ. Батлах гарын үсэг нь тушаал дээр тавигдана.
    gap = _para(doc, "", after=16)
    gap.paragraph_format.keep_with_next = True
    for label, role in (("Боловсруулсан", meta.get("author", HEAD)),
                        ("Хянан баталгаажуулсан", meta.get("reviewer", DIRECTOR))):
        p = _para(doc, "", after=2, line=1.0)
        p.paragraph_format.keep_with_next = True
        _set_font(p.add_run(f"{label}: "), size=11)
        _set_font(p.add_run("……………………………"), size=11)
        p.paragraph_format.tab_stops.add_tab_stop(Cm(9.0))
        _set_font(p.add_run(f"\t/{role}/"), size=11)
    gap2 = _para(doc, "", after=14)
    gap2.paragraph_format.keep_with_next = True

    # ── Хувилбарын түүх ──
    # IT-01 7.3 «хувилбарын бүртгэлд тусгана» гэж шаарддаг ч тийм
    # хэсэг аль ч баримтад байсангүй — заалт хоосон газар зааж байв.
    vh = _para(doc, "Хувилбарын түүх", size=12, bold=True,
               align=WD_ALIGN_PARAGRAPH.CENTER, after=6, line=1.0)
    vh.paragraph_format.keep_with_next = True
    head = ("Хувилбар", "Огноо", "Өөрчлөлтийн товч агуулга",
            "Боловсруулсан", "Баталсан")
    vt = doc.add_table(rows=3, cols=len(head))
    vt.style = "Table Grid"
    widths = (2.0, 2.4, 6.0, 2.6, 2.5)
    rows = [head,
            (meta.get("version", "1.0"), "", "Анхны хувилбар", "", ""),
            ("", "", "", "", "")]
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = vt.cell(i, j)
            cell.width = Cm(widths[j])
            cp = cell.paragraphs[0]
            cp.paragraph_format.space_after = Pt(2)
            cp.paragraph_format.space_before = Pt(2)
            cp.paragraph_format.line_spacing = 1.0
            if i == 0:
                cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _set_font(cp.add_run(val), size=9, bold=(i == 0))
            if i < len(rows) - 1:
                cp.paragraph_format.keep_with_next = True
    for row in vt.rows:
        trPr = row._tr.get_or_add_trPr()
        cant = OxmlElement("w:cantSplit")
        trPr.append(cant)

    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_font(fp.add_run(f"{meta['no']} · {COMPANY_FULL} · {DEPARTMENT} · "),
              size=9, italic=True)
    # Хуудасны дугаар: салангид хуудас дараалал, бүрэн бүтэн байдлаа
    # алддаг тул албан баримтад заавал байна.
    for instr, fallback in ((" PAGE ", "1"), (None, " / "), (" NUMPAGES ", "1")):
        if instr is None:
            _set_font(fp.add_run(fallback), size=9, italic=True)
            continue
        fld = OxmlElement("w:fldSimple")
        fld.set(qn("w:instr"), instr)
        r = OxmlElement("w:r")
        t = OxmlElement("w:t")
        t.text = fallback
        r.append(t)
        fld.append(r)
        fp._p.append(fld)

    # ── Файлын мэдээлэл ──
    # Анхдагчаараа «python-docx» гэж бичигдэж, баримт хэрхэн үүссэнийг
    # ил болгодог. Албан баримтад тохирохгүй.
    cp = doc.core_properties
    cp.title = meta["title"]
    cp.subject = f"{meta['no']} · {DEPARTMENT}"
    cp.author = COMPANY_FULL
    cp.last_modified_by = DEPARTMENT
    cp.category = meta.get("kind", "Журам")
    cp.comments = ""

    path = os.path.join(OUT_DIR, filename)
    doc.save(path)
    _postprocess(path)
    return path


# Хэлний тэмдэглэгээ. Анхдагчаараа en-US байдаг тул Word нээхэд монгол
# үг бүрийг алдаатай гэж улаанаар зурдаг. Албан баримт ингэж харагдах
# ёсгүй бөгөөд уншигчийн анхаарлыг агуулгаас сарниулна.
LANG_RPR = ('<w:lang w:val="mn-MN" w:eastAsia="mn-MN" w:bidi="ar-SA"/>')


def _postprocess(path):
    """Хэлийг mn-MN болгож, үүсгэгч загварын үлдэгдлийг цэвэрлэнэ."""
    tmp = path + ".tmp"
    drop = ("docProps/thumbnail.jpeg", "customXml/item1.xml",
            "customXml/_rels/item1.xml.rels", "customXml/itemProps1.xml")

    with zipfile.ZipFile(path) as zin, \
            zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            name = item.filename
            if name in drop:
                continue
            data = zin.read(name)
            if name in ("word/styles.xml", "word/settings.xml"):
                text = data.decode("utf-8")
                text = re.sub(r'<w:lang [^/]*/>', LANG_RPR, text)
                text = text.replace('w:val="en-US" w:eastAsia="ja-JP"',
                                    'w:val="mn-MN" w:eastAsia="mn-MN"')
                data = text.encode("utf-8")
            elif name in ("[Content_Types].xml", "_rels/.rels",
                          "word/_rels/document.xml.rels"):
                text = data.decode("utf-8")
                for d in drop:
                    text = re.sub(rf'<[^>]*{re.escape(d.split("/")[-1])}[^>]*/>',
                                  "", text)
                text = re.sub(r'<Override[^>]*thumbnail[^>]*/>', "", text)
                data = text.encode("utf-8")
            zout.writestr(item, data)
    shutil.move(tmp, path)
